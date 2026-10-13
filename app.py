import io
import json
import random
import re
import unicodedata
from datetime import datetime
from typing import Dict, List, Optional, Tuple

import streamlit as st

try:
    from pypdf import PdfReader
except ImportError:  # A dependência será instalada pelo requirements.txt.
    PdfReader = None

try:
    from docx import Document
except ImportError:  # A dependência será instalada pelo requirements.txt.
    Document = None


# A chamada deve ser a primeira instrução Streamlit da aplicação.
st.set_page_config(
    page_title="JobMatch AI — Diagnóstico de Currículo",
    page_icon="🎯",
    layout="wide",
    initial_sidebar_state="collapsed",
    menu_items={"About": "# JobMatch AI\nProtótipo acadêmico liderado por Felipe Araújo."},
)

APP_VERSION = "2.0"
MAX_TEXT_CHARS = 50_000
MAX_FILE_BYTES = 5 * 1024 * 1024

# ============================================================================
# Competências e contextos profissionais
# ============================================================================
SKILLS: Dict[str, List[str]] = {
    "Python": ["python", "django", "flask", "fastapi", "pandas", "numpy"],
    "JavaScript": ["javascript", "java script", "ecmascript"],
    "TypeScript": ["typescript", "type script"],
    "Java": ["java", "spring boot", "spring framework"],
    "C/C++": ["c++", "cpp", "linguagem c", "linguagem c++"],
    "C# / .NET": ["c#", "c sharp", ".net", "dotnet", "asp.net"],
    "PHP": ["php", "laravel", "symfony"],
    "HTML/CSS": ["html", "html5", "css", "css3", "folhas de estilo"],
    "React": ["react", "react.js", "reactjs"],
    "Node.js": ["node.js", "nodejs", "node js", "express.js", "expressjs"],
    "SQL": ["sql", "mysql", "postgresql", "postgres", "sqlite", "oracle", "sql server"],
    "Git": ["git", "github", "gitlab", "bitbucket", "versionamento"],
    "Lógica de Programação": [
        "lógica", "logica", "algoritmo", "algoritmos", "estrutura de dados",
        "programação", "programacao", "complexidade de algoritmos",
    ],
    "APIs / REST": ["api", "apis", "rest", "restful", "http", "json", "endpoint", "web service"],
    "Banco de Dados": ["banco de dados", "database", "modelagem de dados", "modelagem de banco", "nosql", "mongodb", "redis"],
    "Cloud / DevOps": ["aws", "azure", "gcp", "cloud", "devops", "docker", "kubernetes", "ci/cd", "pipeline de entrega"],
    "Linux": ["linux", "bash", "shell", "ubuntu", "terminal", "unix"],
    "Testes / QA": ["teste", "testes", "pytest", "unittest", "qa", "testes automatizados", "quality assurance", "selenium", "cypress"],
    "Dados / BI": ["power bi", "powerbi", "dashboard", "business intelligence", "data studio", "tableau", "visualização de dados", "visualizacao de dados"],
    "Excel": ["excel", "planilha", "planilhas", "tabela dinâmica", "tabela dinamica", "procv", "procx"],
    "Automação": ["automação", "automacao", "script", "rpa", "automação de processos", "automacao de processos"],
    "Indústria 4.0": ["indústria 4.0", "industria 4.0", "iot", "manufatura", "automação industrial", "automacao industrial", "plc", "clp", "scada"],
    "Segurança da Informação": ["segurança da informação", "seguranca da informacao", "cibersegurança", "ciberseguranca", "cybersecurity", "owasp", "pentest", "lgpd"],
    "Redes": ["redes de computadores", "tcp/ip", "dns", "dhcp", "roteamento", "firewall", "infraestrutura de redes"],
    "Metodologias Ágeis": ["scrum", "kanban", "metodologias ágeis", "metodologias ageis", "agile", "sprint", "jira"],
    "UX/UI Design": ["ux", "ui", "figma", "prototipação", "prototipacao", "usabilidade", "design de interface"],
}

ROLE_SKILLS = [
    (
        ["desenvolvedor python", "desenvolvedora python", "python developer"],
        ["Python", "Git", "SQL", "APIs / REST", "Lógica de Programação"],
    ),
    (
        ["desenvolvedor web", "desenvolvedora web", "front-end", "frontend", "desenvolvedor front", "web developer"],
        ["HTML/CSS", "JavaScript", "TypeScript", "React", "Git"],
    ),
    (
        ["back-end", "backend", "desenvolvedor backend", "desenvolvedor back-end"],
        ["Python", "Java", "Node.js", "APIs / REST", "SQL", "Git"],
    ),
    (
        ["desenvolvedor", "desenvolvedora", "developer", "programador", "programadora", "desenvolvimento de software", "analista de sistemas"],
        ["Python", "Git", "SQL", "APIs / REST", "Lógica de Programação"],
    ),
    (
        ["analista de dados", "data analyst", "business intelligence", "analista de bi", "bi analyst"],
        ["SQL", "Python", "Dados / BI", "Banco de Dados", "Excel"],
    ),
    (
        ["qa", "quality assurance", "analista de testes", "analista de qualidade de software", "testador", "testadora"],
        ["Testes / QA", "Git", "APIs / REST", "Lógica de Programação"],
    ),
    (
        ["devops", "engenheiro devops", "engenheira devops", "sre", "infraestrutura"],
        ["Git", "Cloud / DevOps", "Linux", "Python", "Redes"],
    ),
    (
        ["suporte de ti", "suporte técnico", "suporte tecnico", "help desk", "service desk", "analista de suporte"],
        ["Linux", "Redes", "Banco de Dados", "Segurança da Informação"],
    ),
    (
        ["cibersegurança", "ciberseguranca", "segurança da informação", "seguranca da informacao", "analista de segurança", "analista de seguranca"],
        ["Segurança da Informação", "Redes", "Linux"],
    ),
    (
        ["automação industrial", "automacao industrial", "indústria", "industria", "manutenção industrial", "manutencao industrial", "instrumentação"],
        ["Automação", "Indústria 4.0", "Python", "Redes"],
    ),
    (
        ["ux designer", "ui designer", "designer de produto", "product designer"],
        ["UX/UI Design", "HTML/CSS", "Metodologias Ágeis"],
    ),
]

SUPPORTED_ROLES = [
    "Desenvolvedor Python Júnior",
    "Desenvolvedor Web Júnior",
    "Desenvolvedor Backend Júnior",
    "Analista de Sistemas Júnior",
    "Analista de Dados Júnior",
    "Analista de BI Júnior",
    "Analista de QA / Testes",
    "Analista de Suporte Técnico",
    "Analista DevOps Júnior",
    "Analista de Segurança da Informação",
    "Analista de Redes",
    "Analista de Automação",
    "Analista de Indústria 4.0",
    "Analista de Banco de Dados",
    "Desenvolvedor Java Júnior",
    "Desenvolvedor JavaScript Júnior",
    "UX/UI Designer Júnior",
]

QUESTION_BANK: Dict[str, List[str]] = {
    "Python": [
        "Como você organizaria um código Python para processar dados de uma API com tratamento de erros?",
        "Qual é a diferença entre lista, tupla e dicionário, e quando escolheria cada estrutura?",
        "Como você validaria os dados recebidos por uma aplicação antes de processá-los?",
        "Como investigaria um erro que acontece em produção, mas não ocorre na sua máquina?",
        "Como você escreveria testes para uma função Python que calcula um desconto?",
    ],
    "JavaScript": [
        "Qual é a diferença entre var, let e const? Em que situação usaria cada uma?",
        "Como você trataria uma chamada assíncrona a uma API usando async/await?",
        "O que é a propagação de eventos no navegador e quando ela é útil?",
        "Como você investigaria um erro que aparece no console do navegador?",
    ],
    "TypeScript": [
        "Quais vantagens o sistema de tipos do TypeScript oferece em relação ao JavaScript puro?",
        "Qual é a diferença entre interface e type no TypeScript?",
        "Como você tiparia a resposta de uma API e trataria campos opcionais?",
    ],
    "Java": [
        "Qual é a diferença entre uma classe abstrata e uma interface em Java?",
        "Como você organizaria classes para representar clientes e pedidos em um sistema?",
        "Como você trataria exceções ao acessar um arquivo ou serviço externo?",
        "O que são encapsulamento e polimorfismo? Dê um exemplo prático.",
    ],
    "C/C++": [
        "Qual é a diferença entre alocação estática e dinâmica de memória?",
        "Como você percorreria um vetor para encontrar o maior valor sem ordenar todos os elementos?",
        "Que cuidados são importantes ao trabalhar com ponteiros em C/C++?",
    ],
    "C# / .NET": [
        "Como você organizaria uma aplicação C# usando classes e interfaces?",
        "Para que serve o mecanismo async/await no .NET?",
        "Como você trataria exceções e registraria erros em uma aplicação .NET?",
    ],
    "PHP": [
        "Como você validaria os dados enviados por um formulário antes de gravá-los?",
        "Como uma aplicação PHP pode acessar um banco de dados com consultas parametrizadas?",
        "Quais cuidados você tomaria para evitar falhas de segurança em uma aplicação PHP?",
    ],
    "HTML/CSS": [
        "Qual é a diferença entre HTML semântico e elementos genéricos como div e span?",
        "Como você faria um layout que se adapta a telas de celular e computador?",
        "Quando usaria Flexbox e quando usaria CSS Grid?",
        "Como você melhoraria a acessibilidade de um formulário web?",
    ],
    "React": [
        "Qual é a diferença entre props e state em React?",
        "Como você carregaria dados de uma API e mostraria estados de carregamento e erro?",
        "Por que as keys são importantes ao renderizar listas em React?",
    ],
    "Node.js": [
        "Como você estruturaria uma rota de API no Node.js com validação e tratamento de erros?",
        "O que significa dizer que o Node.js usa um modelo orientado a eventos?",
        "Como você protegeria variáveis de ambiente e credenciais em uma aplicação Node.js?",
    ],
    "SQL": [
        "Como você faria uma consulta para identificar os cinco clientes com maior valor total comprado?",
        "Qual é a diferença entre INNER JOIN e LEFT JOIN? Dê um exemplo de cada.",
        "Como você usaria JOIN, GROUP BY e ORDER BY para montar um relatório de vendas?",
        "Como investigaria uma consulta SQL que está demorando muito para executar?",
        "Como evitaria resultados duplicados ao relacionar tabelas?",
    ],
    "Git": [
        "Qual seria seu fluxo de trabalho com Git para criar uma funcionalidade sem afetar a branch principal?",
        "Como você resolveria um conflito de merge e verificaria que nenhuma alteração foi perdida?",
        "Qual é a diferença entre commit, branch, merge e pull request?",
        "O que faria se percebesse que enviou uma credencial por engano para um repositório?",
    ],
    "Lógica de Programação": [
        "Qual é a diferença entre uma estrutura condicional e uma estrutura de repetição?",
        "Como você percorreria uma lista e selecionaria apenas os valores que atendem a uma condição?",
        "Como transformaria um problema do cotidiano em etapas lógicas para um programa?",
        "Como verificaria a correção de um algoritmo antes de implementá-lo?",
    ],
    "APIs / REST": [
        "Qual é a diferença entre GET, POST, PUT e DELETE em uma API REST?",
        "Como você validaria uma requisição antes de gravar os dados em um sistema?",
        "Como lidaria com uma API que retorna erros 500 de forma intermitente?",
        "Quais cuidados tomaria com autenticação e autorização em uma API?",
    ],
    "Banco de Dados": [
        "Qual é a diferença entre chave primária e chave estrangeira?",
        "Como modelaria tabelas para relacionar candidatos, vagas e candidaturas?",
        "O que é normalização de dados e qual problema ela ajuda a evitar?",
        "Como investigaria consultas lentas ou inconsistências em um banco de dados?",
    ],
    "Cloud / DevOps": [
        "Quais fatores avaliaria antes de hospedar uma aplicação em nuvem?",
        "Qual é a diferença entre escalar verticalmente e horizontalmente?",
        "Como containers podem facilitar a implantação de uma aplicação?",
        "Como você estruturaria um pipeline simples de integração e entrega contínuas?",
    ],
    "Linux": [
        "Quais comandos usaria para investigar um processo consumindo muita memória?",
        "Como verificaria espaço em disco, memória e processos ativos em Linux?",
        "Como identificaria o processo que está utilizando uma porta específica?",
        "Como analisaria logs para descobrir por que um serviço parou?",
    ],
    "Testes / QA": [
        "Qual é a diferença entre teste unitário, integração e ponta a ponta?",
        "Como montaria casos de teste para uma função que valida um cadastro?",
        "O que faria quando um teste passa localmente, mas falha no ambiente de integração contínua?",
        "Como priorizaria testes para uma funcionalidade com prazo curto?",
    ],
    "Dados / BI": [
        "Como transformaria uma planilha com dados inconsistentes em um dashboard confiável?",
        "Como escolheria indicadores para acompanhar o desempenho de uma operação?",
        "Como explicaria um indicador complexo para uma pessoa da área de negócio?",
        "Como verificaria se os números de um dashboard correspondem à fonte de dados?",
    ],
    "Excel": [
        "Como você usaria filtros, tabelas dinâmicas e fórmulas para resumir uma base de dados?",
        "Como identificaria e corrigiria valores duplicados ou vazios em uma planilha?",
        "Quando usaria PROCV/PROCX ou uma combinação de ÍNDICE e CORRESP?",
    ],
    "Automação": [
        "Que processo repetitivo você automatizaria primeiro e como mediria o resultado?",
        "Como trataria erros para que uma automação não falhe silenciosamente?",
        "Quais cuidados tomaria ao automatizar um processo que manipula dados pessoais?",
    ],
    "Indústria 4.0": [
        "Como sensores e software podem ajudar a detectar uma falha antes da parada de uma máquina?",
        "Como dados de sensores podem apoiar a manutenção preditiva?",
        "Quais riscos devem ser avaliados ao conectar equipamentos industriais a sistemas digitais?",
    ],
    "Segurança da Informação": [
        "Como você explicaria a diferença entre autenticação e autorização?",
        "Quais medidas ajudam a proteger uma aplicação web contra credenciais expostas?",
        "Como priorizaria a correção de vulnerabilidades identificadas em um sistema?",
    ],
    "Redes": [
        "Qual é a diferença entre DNS, DHCP e endereço IP?",
        "Como você investigaria um computador que está conectado ao Wi-Fi, mas não acessa sites?",
        "Para que serve um firewall e como ele pode ajudar a proteger uma rede?",
    ],
    "Metodologias Ágeis": [
        "Qual é o objetivo de uma sprint e como a equipe avalia o que foi concluído?",
        "Como você lidaria com uma tarefa prioritária que surge no meio de uma sprint?",
        "Qual é a diferença entre Scrum e Kanban em um cenário de equipe?",
    ],
    "UX/UI Design": [
        "Como você validaria se um formulário é fácil de entender para novos usuários?",
        "Qual é a diferença entre usabilidade e identidade visual?",
        "Como você usaria um protótipo para recolher feedback antes de desenvolver a interface?",
    ],
}

CUES: Dict[str, List[str]] = {
    "Python": ["funcao", "tratamento", "erro", "try", "except", "lista", "validacao", "dicionario", "teste"],
    "JavaScript": ["const", "let", "async", "await", "evento", "promise", "api"],
    "TypeScript": ["tipo", "interface", "type", "tipagem", "opcional", "compilacao"],
    "Java": ["classe", "interface", "excecao", "encapsulamento", "polimorfismo", "objeto"],
    "C/C++": ["memoria", "ponteiro", "vetor", "alocacao", "indice"],
    "C# / .NET": ["classe", "interface", "async", "await", "excecao", "dotnet"],
    "PHP": ["validacao", "consulta", "parametrizada", "seguranca", "banco"],
    "HTML/CSS": ["semantico", "flexbox", "grid", "responsivo", "acessibilidade", "media query"],
    "React": ["props", "state", "componente", "hook", "api", "key"],
    "Node.js": ["rota", "validacao", "evento", "ambiente", "credencial", "erro"],
    "SQL": ["select", "join", "group", "sum", "order", "limit", "having", "indice"],
    "Git": ["branch", "commit", "pull", "merge", "rebase", "pull request", "conflito"],
    "Lógica de Programação": ["condicao", "repeticao", "if", "for", "while", "algoritmo", "lista"],
    "APIs / REST": ["endpoint", "http", "json", "status", "validacao", "post", "get", "autenticacao"],
    "Banco de Dados": ["tabela", "chave", "relacionamento", "normalizacao", "sql", "indice"],
    "Cloud / DevOps": ["custo", "escala", "seguranca", "aws", "monitoramento", "container", "pipeline"],
    "Linux": ["top", "ps", "free", "kill", "journalctl", "memoria", "processo", "log"],
    "Testes / QA": ["assert", "pytest", "caso", "entrada", "saida", "teste", "regressao"],
    "Dados / BI": ["indicador", "kpi", "dashboard", "dado", "filtro", "visualizacao", "fonte"],
    "Excel": ["tabela dinamica", "procv", "procx", "filtro", "duplicado", "formula"],
    "Automação": ["script", "processo", "tempo", "erro", "ganho", "python", "automatizar"],
    "Indústria 4.0": ["sensor", "iot", "dados", "alerta", "manutencao", "maquina", "preditiva"],
    "Segurança da Informação": ["autenticacao", "autorizacao", "credencial", "vulnerabilidade", "risco", "protecao"],
    "Redes": ["dns", "dhcp", "ip", "firewall", "roteador", "conexao", "rede"],
    "Metodologias Ágeis": ["sprint", "scrum", "kanban", "prioridade", "equipe", "entrega"],
    "UX/UI Design": ["usuario", "usabilidade", "prototipo", "feedback", "interface", "acessibilidade"],
}


# ============================================================================
# Processamento e regras de negócio
# ============================================================================
def norm(text: str) -> str:
    """Normaliza caixa, espaços e acentos para facilitar a comparação textual."""
    text = (text or "").lower().strip()
    text = "".join(
        char
        for char in unicodedata.normalize("NFD", text)
        if unicodedata.category(char) != "Mn"
    )
    return re.sub(r"\s+", " ", text)


def contains_variant(text: str, variant: str) -> bool:
    """Busca termos completos, evitando, por exemplo, detectar 'git' em 'digital'."""
    normalized_text = norm(text)
    normalized_variant = norm(variant)
    if not normalized_variant:
        return False
    # Espaços de uma frase podem aparecer como um ou mais espaços no texto.
    words = normalized_variant.split()
    escaped = r"\s+".join(re.escape(word) for word in words)
    return bool(re.search(r"(?<![a-z0-9])" + escaped + r"(?![a-z0-9])", normalized_text))


def unique(items: List[str]) -> List[str]:
    return list(dict.fromkeys(items))


def extract_skills(text: str) -> List[str]:
    return [
        skill
        for skill, variants in SKILLS.items()
        if any(contains_variant(text, variant) for variant in variants)
    ]


def infer_role_skills(text: str) -> List[str]:
    inferred: List[str] = []
    for triggers, skills in ROLE_SKILLS:
        if any(contains_variant(text, trigger) for trigger in triggers):
            inferred.extend(skill for skill in skills if skill not in inferred)
    return inferred


def analyze(
    resume: str,
    job: str,
    selected_role: str = "Detectar automaticamente pelo anúncio",
    include_suggestions: bool = True,
) -> dict:
    resume_explicit = extract_skills(resume)
    job_explicit = extract_skills(job)
    role_context = "" if selected_role == "Detectar automaticamente pelo anúncio" else selected_role
    job_inferred = [
        skill for skill in infer_role_skills(job + " " + role_context)
        if skill not in job_explicit
    ]
    resume_inferred = [skill for skill in infer_role_skills(resume) if skill not in resume_explicit]
    job_skills = unique(job_explicit + (job_inferred if include_suggestions else []))
    matched = [skill for skill in job_skills if skill in resume_explicit]
    gaps = [skill for skill in job_skills if skill not in resume_explicit]
    score = round((len(matched) / len(job_skills)) * 100) if job_skills else None

    question_skills = unique(job_skills or job_explicit or job_inferred)
    if not question_skills:
        question_skills = unique(resume_explicit)[:5]
    if not question_skills:
        question_skills = ["Lógica de Programação", "Python", "SQL"]

    matrix = []
    for skill in unique(job_explicit + job_inferred):
        if skill in matched:
            status = "Encontrada no currículo"
            priority = "—"
        elif skill in job_explicit:
            status = "Lacuna — requisito escrito na vaga"
            priority = "Alta"
        elif include_suggestions:
            status = "Lacuna — sugestão incluída no cálculo"
            priority = "Média"
        else:
            status = "Sugestão não incluída no cálculo"
            priority = "—"
        matrix.append({"Competência": skill, "Origem": "Vaga" if skill in job_explicit else "Contexto do cargo", "Status": status, "Prioridade": priority})

    return {
        "resume_explicit": resume_explicit,
        "resume_inferred": resume_inferred,
        "job_explicit": job_explicit,
        "job_inferred": job_inferred,
        "job_skills": job_skills,
        "matched": matched,
        "gaps": gaps,
        "score": score,
        "question_skills": question_skills,
        "skill_matrix": matrix,
        "include_suggestions": include_suggestions,
        "selected_role": role_context or "Detecção automática",
    }


def generate_questions(
    skills: List[str],
    amount: int = 3,
    excluded_questions: Optional[List[str]] = None,
) -> List[Tuple[str, str]]:
    """Sorteia perguntas relevantes, evitando repetição na rodada e, quando possível, entre rodadas."""
    amount = max(1, min(int(amount), 8))
    available = [skill for skill in unique(skills) if skill in QUESTION_BANK]
    if not available:
        available = ["Lógica de Programação", "Python", "SQL"]

    excluded = set(excluded_questions or [])
    all_pairs = [(skill, q) for skill in available for q in QUESTION_BANK[skill]]
    fresh_pairs = [(skill, q) for skill, q in all_pairs if q not in excluded]
    if len(fresh_pairs) < amount:
        fresh_pairs = all_pairs

    # Primeiro, tenta distribuir as perguntas por competências diferentes.
    picked: List[Tuple[str, str]] = []
    shuffled_skills = available[:]
    random.shuffle(shuffled_skills)
    for skill in shuffled_skills:
        if len(picked) >= amount:
            break
        candidates = [(s, q) for s, q in fresh_pairs if s == skill and (s, q) not in picked]
        if candidates:
            picked.append(random.choice(candidates))

    remaining = [(s, q) for s, q in fresh_pairs if (s, q) not in picked]
    random.shuffle(remaining)
    picked.extend(remaining[: max(0, amount - len(picked))])

    # Proteção adicional caso o número solicitado seja maior que as perguntas únicas disponíveis.
    while len(picked) < amount:
        picked.append(random.choice(all_pairs))
    return picked[:amount]


def feedback(answer: str, skill: str) -> Tuple[int, str, str]:
    answer = (answer or "").strip()
    words = re.findall(r"\b[\w+#.-]+\b", norm(answer))
    word_count = len(words)
    if word_count < 8:
        return (
            40,
            "Aprofunde",
            "Sua resposta está muito curta para uma avaliação útil. Explique o raciocínio, descreva os passos e cite um exemplo prático.",
        )

    normalized_answer = norm(answer)
    cues = CUES.get(skill, [])
    hits = sum(1 for cue in cues if contains_variant(normalized_answer, cue))
    length_points = min(15, word_count // 5)
    score = max(35, min(95, 50 + hits * 6 + length_points))

    if score >= 82:
        level = "Muito bom"
        message = "Sua resposta contém vários sinais relacionados ao tema. Para fortalecê-la ainda mais, explique uma decisão técnica e dê um exemplo concreto."
    elif score >= 65:
        level = "Bom começo"
        message = "Você trouxe alguns conceitos relevantes. Organize a resposta em etapas, explique por que escolheria essa abordagem e inclua um exemplo prático."
    else:
        level = "Aprofunde"
        message = "A resposta pode ganhar profundidade. Defina os conceitos principais, descreva como resolveria o problema e mencione como testaria o resultado."
    message += " A pontuação é heurística: considera tamanho e termos relacionados, não comprova por si só a correção técnica."
    return score, level, message


def extract_file_text(uploaded_file) -> str:
    """Extrai texto de TXT, PDF ou DOCX; lança ValueError com mensagem amigável."""
    if uploaded_file is None:
        raise ValueError("Selecione um arquivo antes de carregá-lo.")
    if uploaded_file.size > MAX_FILE_BYTES:
        raise ValueError("O arquivo ultrapassa o limite de 5 MB. Use um arquivo menor.")

    suffix = uploaded_file.name.rsplit(".", 1)[-1].lower() if "." in uploaded_file.name else ""
    uploaded_file.seek(0)
    try:
        if suffix == "txt":
            raw = uploaded_file.getvalue()
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("latin-1")
        elif suffix == "pdf":
            if PdfReader is None:
                raise ValueError("A leitura de PDF não está disponível. Instale as dependências de requirements.txt.")
            reader = PdfReader(io.BytesIO(uploaded_file.getvalue()))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            if not text.strip():
                raise ValueError("Não encontrei texto selecionável neste PDF. Se for um documento escaneado, cole o texto manualmente.")
        elif suffix == "docx":
            if Document is None:
                raise ValueError("A leitura de DOCX não está disponível. Instale as dependências de requirements.txt.")
            document = Document(io.BytesIO(uploaded_file.getvalue()))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
            text = "\n".join(parts)
        else:
            raise ValueError("Formato não suportado. Envie um arquivo PDF, DOCX ou TXT.")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Não foi possível ler o arquivo. Verifique se ele não está corrompido. Detalhe: {exc}") from exc

    text = text.strip()
    if not text:
        raise ValueError("O arquivo não contém texto que possa ser usado na análise.")
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS]
    return text


def build_report(data: dict, questions: List[Tuple[str, str]], answers: dict) -> str:
    timestamp = datetime.now().strftime("%d/%m/%Y %H:%M")
    lines = [
        "JOBMATCH AI — RELATÓRIO DO DIAGNÓSTICO",
        f"Gerado em: {timestamp}",
        f"Cargo selecionado: {data.get('selected_role', 'Detecção automática')}",
        "",
        f"Compatibilidade: {str(data['score']) + '%' if data['score'] is not None else 'Não calculada — nenhuma competência identificada na vaga'}",
        f"Competências identificadas na vaga: {', '.join(data['job_explicit']) or 'Nenhuma'}",
        f"Competências sugeridas pelo contexto: {', '.join(data['job_inferred']) or 'Nenhuma'}",
        f"Competências encontradas no currículo: {', '.join(data['matched']) or 'Nenhuma'}",
        f"Lacunas: {', '.join(data['gaps']) or 'Nenhuma'}",
        "",
        "MATRIZ DE COMPETÊNCIAS",
    ]
    for row in data["skill_matrix"]:
        lines.append(f"- {row['Competência']} | {row['Origem']} | {row['Status']} | Prioridade: {row['Prioridade']}")
    lines.extend(["", "ENTREVISTA TÉCNICA"])
    for i, (skill, question) in enumerate(questions):
        lines.append(f"\nPergunta {i + 1} — {skill}: {question}")
        lines.append(f"Resposta: {answers.get(f'answer_{i}', '').strip() or '(não respondida)'}")
        result = answers.get(f"feedback_{i}")
        if result:
            lines.append(f"Feedback: {result[0]}/100 — {result[1]}. {result[2]}")
    lines.extend([
        "",
        "OBSERVAÇÃO",
        "O indicador de compatibilidade é uma comparação textual de competências e não representa probabilidade de contratação.",
        "O feedback das respostas usa heurísticas e não substitui uma avaliação técnica humana nem uma IA generativa.",
    ])
    return "\n".join(lines)


# ============================================================================
# Interface e aparência
# ============================================================================
def inject_css(dark_mode: bool = False) -> None:
    light_css = """
    <style>
    .stApp { background: linear-gradient(180deg, #f7f9fc 0%, #eef3fb 100%); }
    [data-testid="stHeader"] { background: rgba(247,249,252,.88); }
    .block-container { padding-top: 1.3rem; padding-bottom: 3.5rem; max-width: 1280px; }
    .jm-topbar { display:flex; justify-content:space-between; align-items:center; padding:8px 0 16px; border-bottom:1px solid #e4e7ec; margin-bottom:24px; }
    .jm-brand { font-size:1.25rem; font-weight:900; letter-spacing:-.04em; color:#101828; }
    .jm-brand span { color:#2457ff; }
    .jm-badge { padding:7px 10px; border:1px solid #dbe2f0; border-radius:999px; color:#48536a; background:#fff; font-size:.72rem; font-weight:700; }
    .jm-hero { background:linear-gradient(135deg,#101b38 0%,#172b5d 48%,#2457ff 100%); border-radius:26px; padding:32px 34px; color:white; box-shadow:0 20px 48px rgba(23,43,93,.16); position:relative; overflow:hidden; margin-bottom:24px; }
    .jm-eyebrow { font-size:.72rem; letter-spacing:.13em; text-transform:uppercase; font-weight:800; color:#b7c9ff; }
    .jm-title { font-size:2.45rem; line-height:1.06; font-weight:850; letter-spacing:-.045em; margin:.55rem 0 .7rem; max-width:800px; }
    .jm-subtitle { max-width:820px; color:#d7e0f8; line-height:1.6; margin:0; }
    .jm-hero-meta { display:flex; flex-wrap:wrap; gap:8px; margin-top:18px; }
    .jm-hero-pill { border:1px solid rgba(255,255,255,.2); background:rgba(255,255,255,.08); border-radius:999px; padding:7px 10px; font-size:.72rem; color:#e7edff; }
    .jm-section-title { font-size:1.08rem; font-weight:850; color:#101828; margin:24px 0 6px; }
    .jm-section-copy { color:#667085; font-size:.84rem; line-height:1.55; margin-bottom:12px; }
    .jm-role-wrap { display:flex; flex-wrap:wrap; gap:7px; margin:10px 0 8px; }
    .jm-role { border:1px solid #dfe4ed; background:white; color:#344054; border-radius:999px; padding:7px 10px; font-size:.72rem; font-weight:700; }
    .jm-card-title { font-size:.98rem; font-weight:850; color:#101828; margin-bottom:3px; }
    .jm-card-copy { color:#667085; font-size:.77rem; line-height:1.5; margin-bottom:10px; }
    .jm-number { display:inline-flex; width:26px; height:26px; border-radius:9px; align-items:center; justify-content:center; background:#edf2ff; color:#2457ff; font-size:.72rem; font-weight:900; margin-right:7px; }
    div[data-testid="stVerticalBlockBorderWrapper"] { border-radius:18px; border-color:#e4e7ec; box-shadow:0 10px 30px rgba(16,24,40,.05); background:rgba(255,255,255,.93); }
    .stTextArea textarea { border-radius:13px !important; background:#fbfcfe !important; border-color:#dfe4ed !important; }
    .stTextArea textarea:focus { border-color:#9db2ff !important; box-shadow:0 0 0 3px #edf1ff !important; }
    .stButton > button, .stDownloadButton > button { border-radius:12px; font-weight:800; min-height:42px; }
    div.stButton > button[kind="primary"] { background:#2457ff; border-color:#2457ff; }
    div.stButton > button[kind="secondary"], .stDownloadButton > button { background:#f4f6fa; border-color:#e1e6ef; color:#344054; }
    [data-testid="stMetric"] { background:#fff; border:1px solid #e4e7ec; padding:15px; border-radius:15px; box-shadow:0 8px 24px rgba(16,24,40,.04); }
    [data-testid="stMetricLabel"] p { color:#667085 !important; font-weight:700; }
    [data-testid="stMetricValue"] { color:#101828; }
    .jm-chip-row { display:flex; flex-wrap:wrap; gap:6px; margin-top:7px; }
    .jm-chip { display:inline-block; padding:6px 9px; border-radius:999px; background:#f2f4f7; color:#344054; font-size:.71rem; font-weight:750; margin:2px; }
    .jm-chip.ok { background:#eaf8f1; color:#157347; }
    .jm-chip.miss { background:#fff2ef; color:#b42318; }
    .jm-status { padding:11px 12px; border-radius:12px; margin:7px 0; font-size:.8rem; line-height:1.5; }
    .jm-status.ok { background:#e9f8f0; color:#157347; border:1px solid #cceedd; }
    .jm-status.warn { background:#fff5e8; color:#8a4b00; border:1px solid #f6dfb8; }
    .jm-status.danger { background:#fff0ef; color:#b42318; border:1px solid #f4c7c3; }
    .jm-question { border:1px solid #e4e7ec; background:#fbfcff; padding:13px; border-radius:14px; margin:6px 0 12px; }
    .jm-question-head { display:flex; gap:10px; align-items:flex-start; }
    .jm-qnum { width:29px; height:29px; border-radius:9px; background:#edf2ff; color:#2457ff; display:flex; align-items:center; justify-content:center; font-weight:900; flex:0 0 auto; }
    .jm-focus { color:#2457ff; font-size:.7rem; font-weight:850; text-transform:uppercase; letter-spacing:.07em; }
    .jm-qtext { color:#101828; font-size:.94rem; line-height:1.5; font-weight:750; margin-top:2px; }
    .jm-feedback { padding:12px 13px; border-radius:13px; background:#f4f7fb; border:1px solid #e4e8f0; margin-top:9px; font-size:.8rem; line-height:1.6; color:#344054; }
    .jm-footer { color:#98a2b3; text-align:center; font-size:.7rem; margin-top:28px; }
    @media (max-width: 900px) { .jm-title { font-size:1.9rem; } .jm-hero { padding:24px 20px; } }
    </style>
    """
    dark_css = """
    <style>
    .stApp { background:linear-gradient(180deg,#0a0f1b 0%,#11182a 100%) !important; color:#edf2ff; }
    [data-testid="stHeader"] { background:rgba(10,15,27,.9) !important; }
    [data-testid="stMarkdownContainer"] p, [data-testid="stMarkdownContainer"] li, .stCaption, [data-testid="stWidgetLabel"] p { color:#c0cadc; }
    .jm-topbar { border-bottom-color:#253047 !important; }
    .jm-brand, .jm-section-title, .jm-card-title, .jm-qtext, [data-testid="stMetricValue"] { color:#f3f6ff !important; }
    .jm-section-copy, .jm-card-copy, .jm-footer, [data-testid="stMetricLabel"] p { color:#a8b2c7 !important; }
    .jm-badge { background:#141d30 !important; border-color:#2a3853 !important; color:#dce4f5 !important; }
    .jm-role { background:#171f31 !important; border-color:#2a3853 !important; color:#e1e8f5 !important; }
    div[data-testid="stVerticalBlockBorderWrapper"], [data-testid="stMetric"] { background:#121827 !important; border-color:#253047 !important; box-shadow:0 12px 32px rgba(0,0,0,.22) !important; }
    .stTextArea textarea { background:#0e1626 !important; color:#eef3ff !important; border-color:#2a3853 !important; }
    .stTextArea textarea::placeholder { color:#76839b !important; }
    .jm-question { background:#0f1727 !important; border-color:#29364d !important; }
    .jm-feedback { background:#161f31 !important; border-color:#29364d !important; color:#c7d0e1 !important; }
    .jm-chip { background:#1c2639 !important; color:#d8e0ef !important; }
    .jm-chip.ok, .jm-status.ok { background:#103526 !important; border-color:#19583f !important; color:#75ddb0 !important; }
    .jm-chip.miss, .jm-status.danger { background:#351b20 !important; border-color:#6b2829 !important; color:#ff9b9f !important; }
    .jm-status.warn { background:#352719 !important; border-color:#69471f !important; color:#ffbf75 !important; }
    div.stButton > button[kind="secondary"], .stDownloadButton > button { background:#1b2437 !important; color:#e4eaf7 !important; border-color:#2a3853 !important; }
    [data-testid="stDataFrame"] { border:1px solid #29364d; }
    </style>
    """
    st.markdown(light_css, unsafe_allow_html=True)
    if dark_mode:
        st.markdown(dark_css, unsafe_allow_html=True)


def badge(text: str, cls: str = "") -> str:
    return f'<span class="jm-chip {cls}">{text}</span>'


def render_role_section() -> None:
    roles_html = "".join(f'<span class="jm-role">{role}</span>' for role in SUPPORTED_ROLES)
    st.markdown('<div class="jm-section-title">🎯 Cargos que o JobMatch atende</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="jm-section-copy">Exemplos de perfis de TI, dados, suporte, desenvolvimento e indústria. O diagnóstico também pode ser usado com outras vagas.</div>',
        unsafe_allow_html=True,
    )
    st.markdown(f'<div class="jm-role-wrap">{roles_html}</div>', unsafe_allow_html=True)


def render_analysis(data: dict) -> None:
    st.markdown('<div class="jm-section-title">Seu diagnóstico</div>', unsafe_allow_html=True)
    st.markdown('<div class="jm-section-copy">Use os resultados como orientação textual, não como uma decisão automática de contratação.</div>', unsafe_allow_html=True)

    m1, m2, m3 = st.columns(3, gap="medium")
    m1.metric("Compatibilidade", f"{data['score']}%" if data["score"] is not None else "—")
    m2.metric("Competências consideradas", len(data["job_skills"]))
    m3.metric("Lacunas identificadas", len(data["gaps"]))
    if data["score"] is not None:
        st.progress(data["score"] / 100, text=f"Aderência textual: {data['score']}%")
        st.caption("O percentual é a proporção de competências consideradas na vaga que também foram encontradas explicitamente no currículo.")
    else:
        st.warning("Não foi possível calcular a compatibilidade porque nenhuma competência conhecida foi identificada na descrição. Cole mais requisitos ou escolha um cargo para sugerir competências relacionadas.")

    a, b, c = st.columns(3, gap="medium")
    with a:
        with st.container(border=True):
            st.markdown('<div class="jm-card-title">Competências compatíveis</div>', unsafe_allow_html=True)
            st.markdown('<div class="jm-card-copy">Encontradas no currículo e consideradas na vaga.</div>', unsafe_allow_html=True)
            if data["matched"]:
                st.markdown('<div class="jm-chip-row">' + "".join(badge(x, "ok") for x in data["matched"]) + "</div>", unsafe_allow_html=True)
            else:
                st.markdown('<div class="jm-status danger">Nenhuma competência compatível foi identificada explicitamente.</div>', unsafe_allow_html=True)
            if data["resume_explicit"]:
                st.caption("Competências detectadas no currículo: " + " • ".join(data["resume_explicit"]))
            else:
                st.warning("Adicione habilidades, cursos ou projetos técnicos ao currículo para melhorar a análise.")
            if data["resume_inferred"]:
                st.caption("Contexto profissional percebido, não usado como comprovação: " + " • ".join(data["resume_inferred"]))

    with b:
        with st.container(border=True):
            st.markdown('<div class="jm-card-title">Habilidades consideradas</div>', unsafe_allow_html=True)
            st.markdown('<div class="jm-card-copy">Requisitos escritos na vaga e sugestões relacionadas ao cargo.</div>', unsafe_allow_html=True)
            if data["job_explicit"]:
                st.markdown('<div class="jm-chip-row">' + "".join(badge(x) for x in data["job_explicit"]) + "</div>", unsafe_allow_html=True)
            else:
                st.info("Nenhum requisito técnico explícito detectado no texto da vaga.")
            if data["job_inferred"]:
                st.caption("Sugestões pelo contexto do cargo: " + " • ".join(data["job_inferred"]))
            if not data["job_skills"]:
                st.warning("Aumente a descrição da vaga ou informe o cargo no campo de opções.")

    with c:
        with st.container(border=True):
            st.markdown('<div class="jm-card-title">Lacunas prioritárias</div>', unsafe_allow_html=True)
            st.markdown('<div class="jm-card-copy">Competências que ainda não foram identificadas no texto do currículo.</div>', unsafe_allow_html=True)
            if data["gaps"]:
                high = [skill for skill in data["gaps"] if skill in data["job_explicit"]]
                medium = [skill for skill in data["gaps"] if skill not in data["job_explicit"]]
                for gap in high:
                    st.markdown(f'<div class="jm-status warn"><strong>{gap} · prioridade alta</strong><br>É um requisito reconhecido diretamente no anúncio. Mostre um projeto, curso ou experiência se possuir essa habilidade.</div>', unsafe_allow_html=True)
                for gap in medium:
                    st.markdown(f'<div class="jm-status warn"><strong>{gap} · prioridade média</strong><br>É uma sugestão associada ao contexto do cargo. Confirme se ela é realmente exigida na vaga.</div>', unsafe_allow_html=True)
            elif data["job_skills"]:
                st.markdown('<div class="jm-status ok">Nenhuma lacuna encontrada nas competências consideradas. Isso não garante aderência completa à vaga.</div>', unsafe_allow_html=True)
            else:
                st.info("As lacunas aparecerão assim que forem detectadas competências na vaga.")

    if data["skill_matrix"]:
        with st.expander("Ver matriz completa de competências"):
            st.dataframe(data["skill_matrix"], use_container_width=True, hide_index=True)


def render_interview(data: dict, question_count: int) -> None:
    st.markdown('<div class="jm-section-title">Entrevista técnica personalizada</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="jm-section-copy">As perguntas são sorteadas de um banco organizado por competência. Você pode gerar uma nova rodada sem refazer o diagnóstico.</div>',
        unsafe_allow_html=True,
    )
    left, right = st.columns([1.4, 1], gap="medium")
    with left:
        st.markdown("**Como funciona**")
        st.caption(f"A rodada atual tem {len(st.session_state.interview_questions)} pergunta(s). Responda e solicite feedback; o botão de sorteio gera outras questões relacionadas ao perfil.")
    with right:
        if st.button("↻ Sortear novas perguntas", key="shuffle_questions", use_container_width=True):
            previous = list(st.session_state.get("recent_questions", []))
            current_questions = [q for _, q in st.session_state.get("interview_questions", [])]
            previous.extend(current_questions)
            st.session_state.recent_questions = previous[-80:]
            st.session_state.interview_questions = generate_questions(data["question_skills"], question_count, st.session_state.recent_questions)
            for i in range(8):
                st.session_state.pop(f"feedback_{i}", None)
                st.session_state.pop(f"answer_{i}", None)
            st.rerun()

    for i, (skill, question) in enumerate(st.session_state.interview_questions):
        with st.container(border=True):
            st.markdown(
                f'<div class="jm-question"><div class="jm-question-head"><div class="jm-qnum">{i + 1}</div><div><div class="jm-focus">Foco: {skill}</div><div class="jm-qtext">{question}</div></div></div></div>',
                unsafe_allow_html=True,
            )
            answer = st.text_area(
                "Sua resposta",
                key=f"answer_{i}",
                height=125,
                placeholder="Explique o raciocínio, descreva as etapas e, se possível, inclua um exemplo prático.",
            )
            if st.button(f"Receber feedback — pergunta {i + 1}", key=f"feedback_btn_{i}", use_container_width=True):
                if not answer.strip():
                    st.warning("Digite uma resposta antes de solicitar o feedback.")
                else:
                    st.session_state[f"feedback_{i}"] = feedback(answer, skill)
            if f"feedback_{i}" in st.session_state:
                score_feedback, level, message = st.session_state[f"feedback_{i}"]
                st.markdown(
                    f'<div class="jm-feedback"><strong>{score_feedback}/100 — {level}</strong><br>{message}</div>',
                    unsafe_allow_html=True,
                )


def main() -> None:
    dark_mode = st.session_state.get("dark_mode", False)
    inject_css(dark_mode)

    brand_col, badge_col, theme_col = st.columns([4.4, 3.1, 2.0], vertical_alignment="center")
    with brand_col:
        st.markdown('<div class="jm-topbar"><div class="jm-brand">JobMatch <span>AI</span></div></div>', unsafe_allow_html=True)
    with badge_col:
        st.markdown('<div style="padding-top:10px;text-align:right;"><span class="jm-badge">Líder: Felipe Araújo · Python + Streamlit</span></div>', unsafe_allow_html=True)
    with theme_col:
        st.toggle("🌙 Modo escuro", key="dark_mode", help="Alterne entre os modos claro e escuro.")

    st.markdown(
        '<div class="jm-hero"><div class="jm-eyebrow">Diagnóstico de currículo + simulador de entrevista</div><div class="jm-title">Descubra o que falta para seu perfil chegar mais perto da vaga.</div><div class="jm-subtitle">Compare seu currículo com a descrição de uma vaga, identifique competências e pratique perguntas técnicas em uma experiência guiada.</div><div class="jm-hero-meta"><span class="jm-hero-pill">✓ Compatibilidade textual</span><span class="jm-hero-pill">✓ Lacunas prioritárias</span><span class="jm-hero-pill">✓ Perguntas sorteadas</span><span class="jm-hero-pill">✓ Feedback imediato</span></div></div>',
        unsafe_allow_html=True,
    )
    render_role_section()

    st.markdown('<div class="jm-section-title">Comece seu diagnóstico</div>', unsafe_allow_html=True)
    st.markdown('<div class="jm-section-copy">Cole os textos abaixo ou carregue um currículo em PDF, DOCX ou TXT. O arquivo é usado apenas para preencher o campo de texto desta sessão.</div>', unsafe_allow_html=True)

    # Botões de preenchimento aparecem antes dos campos, para que a atualização de estado seja segura no Streamlit.
    sample_col, clear_col, option_col = st.columns([1.25, 1, 2.5], gap="small")
    with sample_col:
        if st.button("✨ Carregar exemplo", use_container_width=True):
            st.session_state.resume_text = (
                "João Silva\nDesenvolvedor de Software Júnior\n\n"
                "Habilidades: Python, Git, SQL, APIs REST e lógica de programação.\n"
                "Projeto: desenvolvimento de uma aplicação web em Python com banco de dados."
            )
            st.session_state.job_text = (
                "Desenvolvedor Python Júnior\n\nRequisitos: Python, Git, SQL, APIs REST, lógica de programação e testes automatizados.\n"
                "Diferenciais: Docker e trabalho com metodologias ágeis."
            )
            st.session_state.selected_role = "Desenvolvedor Python Júnior"
            st.rerun()
    with clear_col:
        if st.button("Limpar campos", use_container_width=True):
            for key in ["resume_text", "job_text"]:
                st.session_state[key] = ""
            for key in list(st.session_state.keys()):
                if key == "analysis" or key == "interview_questions" or key == "recent_questions" or key.startswith("answer_") or key.startswith("feedback_"):
                    st.session_state.pop(key, None)
            st.rerun()
    with option_col:
        selected_role = st.selectbox(
            "Cargo de referência (opcional)",
            ["Detectar automaticamente pelo anúncio"] + SUPPORTED_ROLES,
            key="selected_role",
            help="Use esta opção se o anúncio não deixar claro o cargo ou tiver poucos requisitos técnicos.",
        )

    left, right = st.columns(2, gap="large")
    with left:
        with st.container(border=True):
            st.markdown('<div class="jm-card-title"><span class="jm-number">01</span>Seu currículo</div>', unsafe_allow_html=True)
            st.markdown('<div class="jm-card-copy">Experiências, projetos, cursos e habilidades tornam o diagnóstico mais útil.</div>', unsafe_allow_html=True)
            uploaded = st.file_uploader("Carregar currículo (PDF, DOCX ou TXT)", type=["pdf", "docx", "txt"], key="resume_upload")
            if st.button("Extrair texto do arquivo para o campo", key="load_resume_file", use_container_width=True):
                try:
                    extracted = extract_file_text(uploaded)
                    st.session_state.resume_text = extracted
                    st.toast("Texto extraído. Confira o conteúdo antes de analisar.", icon="✅")
                    st.rerun()
                except ValueError as exc:
                    st.error(str(exc))
            resume = st.text_area(
                "Texto do currículo",
                height=245,
                key="resume_text",
                placeholder="Ex.: João Silva\nDesenvolvedor de Software Júnior\n\nHabilidades: Python, Git, SQL e APIs REST.\nProjeto de automação de relatórios...",
            )
            st.caption(f"{len(resume):,} caracteres · máximo recomendado: {MAX_TEXT_CHARS:,}".replace(",", "."))

    with right:
        with st.container(border=True):
            st.markdown('<div class="jm-card-title"><span class="jm-number">02</span>Vaga desejada</div>', unsafe_allow_html=True)
            st.markdown('<div class="jm-card-copy">Cole o anúncio completo ou, pelo menos, os requisitos e responsabilidades principais.</div>', unsafe_allow_html=True)
            job = st.text_area(
                "Descrição da vaga",
                height=315,
                key="job_text",
                placeholder="Ex.: Desenvolvedor Python Júnior\n\nRequisitos: Python, Git, SQL, APIs REST e lógica de programação.\nDiferenciais: Docker e metodologias ágeis...",
            )
            st.caption(f"{len(job):,} caracteres · máximo recomendado: {MAX_TEXT_CHARS:,}".replace(",", "."))

    options_col1, options_col2 = st.columns([1.2, 2], gap="medium")
    with options_col1:
        question_count = st.select_slider("Quantidade de perguntas", options=[3, 4, 5, 6, 7, 8], value=3)
    with options_col2:
        include_suggestions = st.checkbox(
            "Incluir competências sugeridas pelo cargo no cálculo",
            value=True,
            help="Quando ativo, competências inferidas pelo contexto do cargo entram na pontuação e são identificadas separadamente.",
        )

    if st.button("Analisar meu perfil →", type="primary", use_container_width=True):
        if not resume.strip() or not job.strip():
            st.error("Preencha o currículo e a descrição da vaga antes de analisar.")
        elif len(resume) > MAX_TEXT_CHARS or len(job) > MAX_TEXT_CHARS:
            st.error(f"Cada campo deve ter no máximo {MAX_TEXT_CHARS:,} caracteres.".replace(",", "."))
        else:
            analysis = analyze(resume, job, selected_role, include_suggestions)
            st.session_state.analysis = analysis
            st.session_state.analysis_input = {"resume": resume, "job": job, "selected_role": selected_role, "question_count": question_count}
            st.session_state.interview_questions = generate_questions(analysis["question_skills"], question_count)
            st.session_state.recent_questions = [q for _, q in st.session_state.interview_questions]
            for i in range(8):
                st.session_state.pop(f"answer_{i}", None)
                st.session_state.pop(f"feedback_{i}", None)
            st.session_state.history = st.session_state.get("history", [])
            st.session_state.history.append({
                "data": datetime.now().strftime("%d/%m/%Y %H:%M"),
                "cargo": selected_role if selected_role != "Detectar automaticamente pelo anúncio" else "Detecção automática",
                "score": analysis["score"],
                "lacunas": len(analysis["gaps"]),
            })
            st.session_state.history = st.session_state.history[-10:]
            st.rerun()

    if "analysis" in st.session_state:
        data = st.session_state.analysis
        st.divider()
        render_analysis(data)
        st.divider()
        render_interview(data, st.session_state.get("analysis_input", {}).get("question_count", question_count))

        answers = {key: value for key, value in st.session_state.items() if key.startswith("answer_") or key.startswith("feedback_")}
        questions = st.session_state.interview_questions
        report = build_report(data, questions, answers)
        download_col1, download_col2 = st.columns(2, gap="medium")
        with download_col1:
            st.download_button(
                "⬇️ Baixar relatório (.txt)",
                data=report,
                file_name="jobmatch_relatorio.txt",
                mime="text/plain;charset=utf-8",
                use_container_width=True,
            )
        with download_col2:
            export_payload = {
                "versao": APP_VERSION,
                "gerado_em": datetime.now().isoformat(timespec="minutes"),
                "diagnostico": data,
                "perguntas": [
                    {
                        "competencia": skill,
                        "pergunta": question,
                        "resposta": st.session_state.get(f"answer_{i}", ""),
                        "feedback": st.session_state.get(f"feedback_{i}"),
                    }
                    for i, (skill, question) in enumerate(questions)
                ],
                "observacao": "Pontuação textual e feedback heurístico; não representam probabilidade de contratação.",
            }
            st.download_button(
                "⬇️ Exportar dados (.json)",
                data=json.dumps(export_payload, ensure_ascii=False, indent=2),
                file_name="jobmatch_relatorio.json",
                mime="application/json",
                use_container_width=True,
            )

    if st.session_state.get("history"):
        with st.expander("Histórico desta sessão (até 10 análises)"):
            st.caption("Este histórico existe apenas durante a sessão atual e não é salvo em um banco de dados.")
            st.dataframe(st.session_state.history, use_container_width=True, hide_index=True)

    st.markdown(
        '<div class="jm-footer">JobMatch AI · Python + Streamlit · análise baseada em palavras-chave e regras · perguntas técnicas sorteadas · feedback heurístico</div>',
        unsafe_allow_html=True,
    )


if __name__ == "__main__":
    main()
